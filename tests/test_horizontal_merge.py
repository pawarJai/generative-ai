"""Side-by-side merge and column combine.

Written against a real artifact. outputs/test_f1-2f.xlsx was produced by
asking for two documents in one file, and its Combined sheet is 98 rows x 12
columns:

    SL no ... Total Quantity   col_0 col_1 col_2 col_3
    (64 non-null, rows 1-64)   (34 non-null, rows 65-98)

Both documents WERE read — the 34 rows are really there. What is wrong is the
axis: export_data stacks, so data-file-2's rows landed BELOW data-file-1's
under their own separate columns, and every cell of each block is blank in the
other block's columns. There was no tool that could put them beside each other.

These tests fix the axis (a merge must come out wider, never taller), the
alignment (columns are assembled, never positionally unioned), and the two
integrations a new file-writing tool has to have: the export backstops must
recognise it, or a successful merge gets "recovered" a second time by
export_data and overwritten with a vertical stack.
"""
import os

import pandas as pd
import pytest

from app.graph import agent as A
from app.graph import tools as T


@pytest.fixture
def outdir(tmp_path, monkeypatch):
    monkeypatch.setattr("app.config.OUTPUT_DIR", str(tmp_path))
    return tmp_path


@pytest.fixture
def two_docs(monkeypatch):
    """Two documents whose tables are the shape the real ones are: named
    columns on the left, and a right-hand document of different length."""
    left = pd.DataFrame({
        "SL no": [1, 2, 3, 4],
        "Item": ["40NB PN10", "50NB PN16", "65NB PN10", "80NB PN16"],
        "Spec No.": ["S1", "S2", "S3", "S4"],
    })
    right = pd.DataFrame({
        "Evaluation Schedules": ["Delivery", "Warranty"],
        "Item/Category": ["Valves", "Valves"],
    })
    frames = {"fid1": (left, [8, 9, 10]), "fid2": (right, [6, 7])}

    def fake_side_frame(file_id, pages):
        if file_id not in frames:
            return None, "no extractable tables", []
        df, used = frames[file_id]
        return df.copy(), None, (list(pages) if pages else used)

    monkeypatch.setattr(T, "_side_frame", fake_side_frame)
    monkeypatch.setattr("app.graph.agent._display_name",
                        lambda fid, stored: {"fid1": "data-file-1.pdf",
                                             "fid2": "data-file-2.pdf"}.get(fid, fid))
    from app import state as app_state
    monkeypatch.setattr(app_state, "FILE_ORIGINAL_NAME",
                        {"fid1": "data-file-1.pdf", "fid2": "data-file-2.pdf"})
    monkeypatch.setattr(app_state, "FILE_ORDER", ["fid1", "fid2"])
    monkeypatch.setattr(app_state, "WRITTEN_EXPORTS", {})
    return frames


def _invoke(**kwargs):
    return T.merge_files_side_by_side.invoke(kwargs)


# --- the axis itself -------------------------------------------------------

def test_merge_is_wider_not_taller(outdir, two_docs):
    """The whole defect in one assertion: 4 rows beside 2 rows is 4 rows of 5
    columns. test_f1-2f.xlsx made it 6 rows of 5 columns instead."""
    _invoke(file1_id="fid1", file1_pages=[8, 9, 10],
            file2_id="fid2", file2_pages=[6, 7], output_filename="m.xlsx")
    df = pd.read_excel(outdir / "m.xlsx")
    assert df.shape == (4, 5)


def test_columns_carry_their_source_suffix(outdir, two_docs):
    _invoke(file1_id="fid1", file2_id="fid2", output_filename="m.xlsx")
    cols = list(pd.read_excel(outdir / "m.xlsx").columns)
    assert cols == ["SL no_f1", "Item_f1", "Spec No._f1",
                    "Evaluation Schedules_f2", "Item/Category_f2"]


def test_row_one_of_each_file_sits_on_the_same_row(outdir, two_docs):
    _invoke(file1_id="fid1", file2_id="fid2", output_filename="m.xlsx")
    row = pd.read_excel(outdir / "m.xlsx").iloc[0]
    assert row["Item_f1"] == "40NB PN10"
    assert row["Evaluation Schedules_f2"] == "Delivery"


def test_padding_of_the_shorter_side_is_stated(outdir, two_docs):
    """A merge of unequal tables leaves blank cells. Saying "4 rows x 5
    columns" and stopping would let a half-empty file read as a complete one."""
    out = _invoke(file1_id="fid1", file2_id="fid2", output_filename="m.xlsx")
    assert "Row counts differ (4 vs 2)" in out
    assert "blank on the 'data-file-2.pdf' side" in out
    assert "matched by position" in out


def test_equal_length_merge_says_nothing_about_padding(outdir, two_docs,
                                                       monkeypatch):
    same = pd.DataFrame({"A": [1, 2, 3, 4], "B": [5, 6, 7, 8]})
    monkeypatch.setattr(T, "_side_frame", lambda fid, pages: (
        (same.copy(), None, [1]) if fid == "fid2"
        else two_docs["fid1"][0].copy().pipe(lambda d: (d, None, [8]))))
    out = _invoke(file1_id="fid1", file2_id="fid2", output_filename="m.xlsx")
    assert "Row counts differ" not in out


# --- provenance and the context band --------------------------------------

def test_provenance_columns_are_not_merged_in(outdir, two_docs, monkeypatch):
    """_source_file/_source_page identify ONE document. Side by side there are
    two, so a single pair of provenance columns would be false for half the
    sheet."""
    left = two_docs["fid1"][0].copy()
    left["_source_file"] = "data-file-1.pdf"
    left["_source_page"] = 8
    monkeypatch.setattr(T, "_side_frame", lambda fid, pages: (
        (left.copy(), None, [8]) if fid == "fid1"
        else (two_docs["fid2"][0].copy(), None, [6])))
    _invoke(file1_id="fid1", file2_id="fid2", output_filename="m.xlsx")
    cols = list(pd.read_excel(outdir / "m.xlsx").columns)
    assert not [c for c in cols if "_source" in c]


def test_no_band_is_written_and_that_is_said(outdir, two_docs):
    out = _invoke(file1_id="fid1", file2_id="fid2", output_filename="m.xlsx")
    assert "No header band was written" in out
    first = pd.read_excel(outdir / "m.xlsx", header=None, nrows=1)
    assert int(first.iloc[0].notna().sum()) >= 2  # header row, not a letterhead


# --- refusing rather than guessing ----------------------------------------

def test_one_document_is_refused_not_guessed(outdir, two_docs):
    """FILE_ORDER[0]/[-1] is exactly the fallback that exported a document
    nobody had mentioned. With one file named, this asks instead."""
    out = _invoke(file1_id="fid1", output_filename="m.xlsx")
    assert "needs TWO different documents" in out
    assert not list(outdir.glob("*.xlsx"))


def test_same_document_twice_is_refused(outdir, two_docs):
    out = _invoke(file1_id="fid1", file2_id="fid1", output_filename="m.xlsx")
    assert "needs TWO different documents" in out


def test_unreadable_side_creates_no_file(outdir, two_docs):
    out = _invoke(file1_id="fid1", file2_id="missing", output_filename="m.xlsx")
    assert "No file was created" in out
    assert not list(outdir.glob("*.xlsx"))


def test_pages_are_taken_from_the_users_own_words(outdir, two_docs, monkeypatch):
    """The model's paraphrase drops page numbers. _resolve_source_specs reads
    the real prompt, which is where "pages 8 9 10 and ... 6 7" survives."""
    from app import state as app_state
    monkeypatch.setattr(app_state, "CURRENT_USER_PROMPT",
                        "merge data file 1 page 8 9 10 and data file 2 page 6 7 "
                        "side by side save as mrg-04.xlsx")
    monkeypatch.setattr(
        "app.graph.agent._resolve_source_specs",
        lambda prompt, fid: [{"file_id": "fid1", "pos": 0, "pages": [8, 9, 10]},
                             {"file_id": "fid2", "pos": 30, "pages": [6, 7]}])
    seen = {}

    def spy(file_id, pages):
        seen[file_id] = pages
        return two_docs[file_id][0].copy(), None, pages or []

    monkeypatch.setattr(T, "_side_frame", spy)
    _invoke(output_filename="mrg-04.xlsx")
    assert seen == {"fid1": [8, 9, 10], "fid2": [6, 7]}


# --- filenames -------------------------------------------------------------

def test_output_name_is_forced_to_xlsx(outdir, two_docs):
    _invoke(file1_id="fid1", file2_id="fid2", output_filename="mrg-04")
    assert (outdir / "mrg-04.xlsx").exists()


def test_output_name_cannot_escape_the_outputs_folder(outdir, two_docs):
    _invoke(file1_id="fid1", file2_id="fid2",
            output_filename="../../escaped.xlsx")
    assert (outdir / "escaped.xlsx").exists()
    assert not (outdir.parent / "escaped.xlsx").exists()


def test_merged_file_can_be_modified_afterwards(outdir, two_docs):
    """modify_export refuses with "no file has been exported" unless the file
    is registered, so "now add the header to that" would fail on a merge."""
    from app import state as app_state
    _invoke(file1_id="fid1", file1_pages=[8, 9],
            file2_id="fid2", output_filename="m.xlsx")
    assert app_state.WRITTEN_EXPORTS["m.xlsx"]["file_id"] == "fid1"
    assert app_state.WRITTEN_EXPORTS["m.xlsx"]["pages"] == [8, 9]


# --- suffix derivation (no hardcoded document names) ----------------------

@pytest.mark.parametrize("name,expected", [
    ("data-file-1.pdf", "f1"),
    ("data-file-2.pdf", "f2"),
    ("tender 14.pdf", "f14"),
    ("Quotation.xlsx", "quotat"),
])
def test_suffix_comes_from_the_files_own_name(name, expected):
    assert T._merge_suffix(name) == expected


def test_colliding_suffixes_are_separated(outdir, two_docs, monkeypatch):
    """Two documents that reduce to the same tag would produce duplicate
    column names and an unreadable file."""
    monkeypatch.setattr("app.graph.agent._display_name",
                        lambda fid, stored: "rfq-1.pdf")
    _invoke(file1_id="fid1", file2_id="fid2", output_filename="m.xlsx")
    cols = list(pd.read_excel(outdir / "m.xlsx").columns)
    assert len(cols) == len(set(cols))
    assert any(c.endswith("_f1-l") for c in cols)
    assert any(c.endswith("_f1-r") for c in cols)


# --- combine_columns -------------------------------------------------------

def _write(path, frames: dict, band=0):
    with pd.ExcelWriter(path) as w:
        for name, df in frames.items():
            df.to_excel(w, sheet_name=name, index=False, startrow=band)
    if band:  # a context letterhead above the header row
        from openpyxl import load_workbook
        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.cell(row=1, column=1, value="TENDER 14 — Cochin Shipyard")
        wb.save(path)


def _combine(**kwargs):
    return T.combine_columns.invoke(kwargs)


@pytest.fixture
def sheet(outdir):
    _write(outdir / "mrg-04.xlsx", {"Sheet1": pd.DataFrame({
        "SL no_f1": [1, 2, 3],
        "Item_f1": ["40NB PN10", "50NB PN16", "65NB PN10"],
        "Evaluation Schedules_f2": ["Delivery", "Warranty", None],
    })})
    from app import state as app_state
    app_state.WRITTEN_EXPORTS.setdefault("mrg-04.xlsx", {"file_id": "fid1"})
    return outdir / "mrg-04.xlsx"


def test_combine_creates_the_new_column(outdir, sheet):
    _combine(source_filename="mrg-04.xlsx",
             columns_to_combine=["SL no_f1", "Item_f1"],
             new_column_name="item_code", separator="-",
             output_filename="mrg-04-combined.xlsx")
    df = pd.read_excel(outdir / "mrg-04-combined.xlsx")
    assert list(df["item_code"]) == ["1-40NB PN10", "2-50NB PN16", "3-65NB PN10"]


def test_combine_skips_blanks_instead_of_writing_nan(outdir, sheet):
    _combine(source_filename="mrg-04.xlsx",
             columns_to_combine=["Item_f1", "Evaluation Schedules_f2"],
             new_column_name="full_description")
    df = pd.read_excel(outdir / "mrg-04_combined.xlsx")
    assert df["full_description"].iloc[2] == "65NB PN10"
    assert "nan" not in " ".join(df["full_description"].astype(str))


def test_combine_matches_column_names_loosely(outdir, sheet):
    """The model retypes 'SL no_f1' as 'SL_no_f1'. An exact comparison would
    reject it and send the user back to copy the name character by character."""
    out = _combine(source_filename="mrg-04.xlsx",
                   columns_to_combine=["SL_no_f1", "item_f1"],
                   new_column_name="item_code")
    assert "Verified" in out
    assert "item_code" in pd.read_excel(outdir / "mrg-04_combined.xlsx").columns


def test_combine_reads_past_a_context_band(outdir):
    """An export written with the header band has its real column names on
    row 2. Reading row 1 as the header reports every column as missing."""
    _write(outdir / "banded.xlsx",
           {"Sheet1": pd.DataFrame({"A": [1, 2], "B": ["x", "y"]})}, band=1)
    out = _combine(source_filename="banded.xlsx",
                   columns_to_combine=["A", "B"], new_column_name="AB")
    assert "Verified" in out
    assert list(pd.read_excel(outdir / "banded_combined.xlsx")["AB"]) == ["1 x", "2 y"]


def test_combine_keeps_the_other_sheets(outdir):
    """A multi-document export is one sheet per file. Rewriting only the first
    would delete the rest."""
    _write(outdir / "multi.xlsx", {
        "one": pd.DataFrame({"A": [1], "B": [2]}),
        "two": pd.DataFrame({"C": [3], "D": [4]}),
    })
    _combine(source_filename="multi.xlsx", columns_to_combine=["C", "D"],
             new_column_name="CD")
    book = pd.ExcelFile(outdir / "multi_combined.xlsx")
    assert book.sheet_names == ["one", "two"]
    assert "CD" in pd.read_excel(outdir / "multi_combined.xlsx", sheet_name="two")


def test_combine_names_the_real_columns_when_it_cannot_match(outdir, sheet):
    out = _combine(source_filename="mrg-04.xlsx",
                   columns_to_combine=["first_name", "last_name"],
                   new_column_name="full_name")
    assert "not in mrg-04.xlsx" in out
    assert "SL no_f1" in out
    assert not (outdir / "mrg-04_combined.xlsx").exists()


def test_combine_on_a_missing_file_lists_what_exists(outdir, sheet):
    out = _combine(source_filename="nope.xlsx",
                   columns_to_combine=["A", "B"], new_column_name="AB")
    assert "not in the outputs folder" in out
    assert "mrg-04.xlsx" in out


def test_combine_needs_two_columns(outdir, sheet):
    out = _combine(source_filename="mrg-04.xlsx",
                   columns_to_combine=["Item_f1"], new_column_name="x")
    assert "at least two" in out


def test_combine_source_is_confined_to_outputs(outdir, sheet):
    out = _combine(source_filename="../../../etc/passwd",
                   columns_to_combine=["A", "B"], new_column_name="AB")
    assert "not in the outputs folder" in out


# --- registration ----------------------------------------------------------

def test_both_tools_are_registered():
    names = {t.name for t in A.TOOLS}
    assert {"merge_files_side_by_side", "combine_columns"} <= names


def test_both_tools_count_as_exports():
    """Left out of this set, a verified merge looks to the backstops like a
    turn where no file was written — and export_data reruns the request as a
    vertical stack, overwriting the merge."""
    assert {"merge_files_side_by_side", "combine_columns"} <= A._EXPORT_TOOL_NAMES
    assert A._tool_to_intent("merge_files_side_by_side") == "export"
    assert A._tool_to_intent("combine_columns") == "export"


# --- the invented-column guard must not fire on a column being created ----

@pytest.mark.parametrize("prompt", [
    # The live turn that produced the false banner.
    "in mrg-04.xlsx combine SL_no_f1 and Item_f1 columns with a dash separator "
    "to create a column called item_code save as mrg-04-code.xlsx",
    "combine the Item_f1 and Evaluation Schedules_f2 columns into one column "
    "called full_description",
    "add a new column joining first_name and last_name",
])
def test_creating_a_column_is_not_a_fabricated_column(prompt, monkeypatch):
    """The guard exists to stop the model listing values for a column the
    document does not have. A column the user asked to BUILD is absent from
    the document by definition — reporting that as fabrication told the user
    to ignore a file that had just been written correctly."""
    monkeypatch.setattr(A, "_header_only_column_names", lambda fid: set())
    out = A._catch_invented_column("Created it.", prompt, "fid1")
    assert out == "Created it."


def test_asking_about_a_real_missing_column_is_still_caught(monkeypatch):
    """The protection itself must survive the fix above."""
    monkeypatch.setattr("app.tables.helpers.get_all_real_tables",
                        lambda fid, **kw: [pd.DataFrame({"Item": [1]})])
    monkeypatch.setattr(A, "_header_only_column_names", lambda fid: set())
    from app import state as app_state
    monkeypatch.setitem(app_state.FILE_KIND, "fid1", "docling")
    out = A._catch_invented_column(
        "Schedule 1, Schedule 2.",
        "give me unique data of column EvaluationSchedules", "fid1")
    assert "there is no column named `EvaluationSchedules`" in out


def test_column_name_capture_stops_at_the_filename(monkeypatch):
    """'column called item_code save as mrg-04-code.xlsx' captured the whole
    tail as the column name, so even the banner's subject was wrong."""
    m = A._COLUMN_ASK_RE.search("create a column called item_code save as x.xlsx")
    assert (m.group(1) or m.group(2)) == "item_code"


def test_existing_tools_survived_the_addition():
    names = {t.name for t in A.TOOLS}
    assert {"search_documents", "query_table_data", "list_uploaded_files",
            "export_data", "modify_export", "get_file_overview",
            "get_page_content", "get_table_of_contents", "generate_quotation",
            "analyze_past_contracts"} <= names
