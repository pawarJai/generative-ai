"""Deterministic export column-selection / row-filter / top-N / computed-
column support.

Confirmed production failure (interaction_log.jsonl, session 477b296a,
2026-08-13 19:32): "in uploaded data file 3 we need to export two column
data in f1-exp-03.xlsx, this column we need to export = Item Title, Item
Quantity, do it" wrote all 7 real columns to the output file, and a
follow-up complaint ("still you add all columns data why") was answered with
a fabricated "verified" reassurance instead of a real fix. Confirmed on disk
via openpyxl: outputs/f1-exp-03.xlsx and f1-exp-04.xlsx both had every
column. export_data's fallback path built the exported table from
run_code_on_files (an LLM sandbox) and handed it to the exporter with no
column filter ever applied.

app.export.spec gives export_data (and app.export.modify, for files that
already exist) a way to resolve column selection, row filtering, row-count
limits, and simple computed columns directly against a table's real columns,
without depending on the sandbox to get it right.
"""
import os

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import state
from app.config import OUTPUT_DIR
from app.export import spec
from app.export.exporters import export_excel
from app.export.modify import modify
from app.models import QueryPlan


REAL_CSV = "uploads/data-file-3_7044_data-file-3.csv"


def _frame():
    return pd.DataFrame({
        "Item Number": list(range(1, 11)),
        "Item Title": [f"Item {i}" for i in range(1, 11)],
        "Item Quantity": [i * 6 for i in range(1, 11)],
        "Unit of Measure": ["EA"] * 10,
    })


# --- extract_requested_columns -------------------------------------------

def test_a_named_subset_of_columns_is_extracted():
    cols = spec.extract_requested_columns(
        "export Item Title, Item Quantity into f.xlsx",
        ["Item Number", "Item Title", "Item Quantity", "Unit of Measure"])
    assert cols == ["Item Title", "Item Quantity"]


def test_naming_every_column_is_not_treated_as_a_selection():
    """If every real column happens to appear in the prompt, that is not the
    same as the user asking to narrow to a subset."""
    all_cols = ["Item Title", "Item Quantity"]
    cols = spec.extract_requested_columns(
        "export Item Title and Item Quantity, all columns", all_cols)
    assert cols == []


def test_no_column_names_in_the_prompt_extracts_nothing():
    cols = spec.extract_requested_columns(
        "export the working sheet to excel",
        ["Item Number", "Item Title", "Item Quantity"])
    assert cols == []


def test_short_column_names_are_not_fuzzy_matched_by_accident():
    """A 2-character column name could match almost anywhere; the >=3
    length guard exists so 'ID' doesn't spuriously match unrelated text."""
    cols = spec.extract_requested_columns(
        "export everything to a file, thanks", ["ID", "Item Title"])
    assert "ID" not in cols


# --- extract_row_filter ----------------------------------------------------

def test_a_where_clause_resolves_to_the_real_column():
    filt = spec.extract_row_filter(
        "export rows where Item Number = 5",
        ["Item Number", "Item Title", "Item Quantity"])
    assert filt == ("Item Number", "5")


def test_a_filter_on_an_unknown_column_is_not_honoured():
    filt = spec.extract_row_filter(
        "export rows where Made Up Column = xyz",
        ["Item Number", "Item Title"])
    assert filt is None


def test_apply_row_filter_matches_case_insensitively():
    df = _frame()
    out = spec.apply_row_filter(df, ("Item Title", "item 5"))
    assert len(out) == 1
    assert out.iloc[0]["Item Title"] == "Item 5"


# --- extract_row_limit ------------------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("give me the top 10 rows", ("head", 10)),
    ("export first five rows only", ("head", 5)),
    ("give me the bottom 5 rows", ("tail", 5)),
    ("export last ten rows", ("tail", 10)),
    ("export everything", None),
])
def test_row_limit_phrasings(text, expected):
    assert spec.extract_row_limit(text) == expected


def test_apply_row_limit_head_and_tail():
    df = _frame()
    assert len(spec.apply_row_limit(df, ("head", 3))) == 3
    assert spec.apply_row_limit(df, ("tail", 3))["Item Number"].tolist() == [8, 9, 10]


# --- extract_computed_column -------------------------------------------------

def test_computed_column_two_real_columns():
    df = _frame()
    computed = spec.extract_computed_column(
        "add a column Total = Item Quantity * Item Number", list(df.columns))
    assert computed == {"name": "Total", "left": "Item Quantity", "op": "*",
                        "right_col": "Item Number", "right_val": None}


def test_computed_column_with_a_literal_number():
    df = _frame()
    computed = spec.extract_computed_column(
        "add a column Doubled = Item Quantity * 2", list(df.columns))
    assert computed["right_col"] is None
    assert computed["right_val"] == 2.0
    out = spec.apply_computed_column(df, computed)
    assert out["Doubled"].tolist() == [q * 2 for q in df["Item Quantity"]]


def test_computed_column_with_an_unresolvable_operand_is_not_returned():
    df = _frame()
    assert spec.extract_computed_column(
        "add a column Total = NotAColumn * 2", list(df.columns)) is None


# --- parse_and_apply: the masking regression --------------------------------

def test_a_filter_clause_does_not_leak_into_column_selection():
    """The real bug caught while building this: 'where Item Title =
    jay-pawar' matched 'Item Title' as if the user had also asked to select
    only that column, narrowing a row filter into an unwanted column filter
    too. The filter clause must be masked out before column selection runs."""
    df = _frame()
    out, changes = spec.parse_and_apply(df, "export rows where Item Title = item 3")
    assert list(out.columns) == list(df.columns)
    assert any("filtered" in c for c in changes)
    assert not any("kept only columns" in c for c in changes)


def test_combining_a_limit_and_a_column_selection():
    df = _frame()
    out, changes = spec.parse_and_apply(
        df, "export bottom 3 rows, only Item Title, Item Quantity columns")
    assert list(out.columns) == ["Item Title", "Item Quantity"]
    assert len(out) == 3


# --- regression: the exact production failure -------------------------------

@pytest.mark.skipif(not os.path.exists(REAL_CSV), reason="real upload not present")
def test_the_exact_failing_export_now_selects_exactly_two_columns():
    df = pd.read_csv(REAL_CSV)
    prompt = ("in uploaded data file 3  we need to export two column data in "
             "f1-exp-03.xlsx  , this column we need to export= Item Title, "
             "Item Quantity  , do it")
    out, changes = spec.parse_and_apply(df, prompt)
    assert list(out.columns) == ["Item Title", "Item Quantity"]
    assert len(out) == 64
    assert any("kept only columns" in c for c in changes)


# --- app.export.modify: the same capabilities on an already-exported file --

def _write(name, df):
    plan = QueryPlan(intent="export", sink="excel", filename=name, no_context=True)
    export_excel("test_file", plan, tables=[df])
    return os.path.join(OUTPUT_DIR, name)


def test_modify_can_filter_rows_of_an_existing_export():
    path = _write("t-spec-filter.xlsx", _frame())
    result = modify("t-spec-filter.xlsx", "keep only rows where Item Number = 5")
    assert "✓" in result
    remaining = pd.read_excel(path)
    assert len(remaining) == 1
    assert remaining.iloc[0]["Item Number"] == 5


def test_modify_can_keep_only_the_top_n_rows():
    path = _write("t-spec-topn.xlsx", _frame())
    result = modify("t-spec-topn.xlsx", "give me only the top 4 rows")
    assert "✓" in result
    assert len(pd.read_excel(path)) == 4


def test_modify_can_add_a_computed_column():
    path = _write("t-spec-calc.xlsx", _frame())
    result = modify("t-spec-calc.xlsx", "add a column Total = Item Quantity * 2")
    assert "✓" in result
    out = pd.read_excel(path)
    assert "Total" in out.columns
    assert out["Total"].tolist() == [q * 2 for q in out["Item Quantity"]]


def test_modify_reports_a_filter_matching_nothing_instead_of_writing_empty():
    path = _write("t-spec-empty.xlsx", _frame())
    before = load_workbook(path).active.max_row
    result = modify("t-spec-empty.xlsx", "keep only rows where Item Title = nonexistent")
    assert "matched no rows" in result
    assert load_workbook(path).active.max_row == before


def test_modify_guidance_message_mentions_the_new_capabilities():
    path = _write("t-spec-guidance.xlsx", _frame())
    result = modify("t-spec-guidance.xlsx", "make this file better somehow")
    assert "top/bottom N rows" in result
    assert "computed column" in result


# --- app.graph.tools.export_data: end-to-end through the real tool --------

def test_export_data_selects_exact_columns_via_the_deterministic_path(monkeypatch, tmp_path):
    from app import state as app_state
    from app.graph.tools import export_data

    table = _frame()
    monkeypatch.setattr("app.tables.helpers.get_all_real_tables",
                        lambda fid, **kw: [table] if fid == "df3" else [])
    monkeypatch.setitem(app_state.FILE_KIND, "df3", "tabular")
    monkeypatch.setattr(app_state, "FILE_ORIGINAL_NAME", {"df3": "data-file-3.csv"})
    monkeypatch.setattr(app_state, "FILE_ORDER", ["df3"])
    monkeypatch.setattr(app_state, "FILE_META", {})
    monkeypatch.setattr(app_state, "CURRENT_USER_PROMPT", None)

    out_name = "t-tool-cols.xlsx"
    result = export_data.invoke({
        "question": "export Item Title, Item Quantity columns",
        "format": "excel",
        "filename": out_name,
        "file_id": "df3",
    })
    assert "Applied" in result
    assert "kept only columns" in result
    written = pd.read_excel(os.path.join(OUTPUT_DIR, out_name))
    assert list(written.columns) == ["Item Title", "Item Quantity"]
    assert len(written) == 10
