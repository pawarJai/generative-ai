"""Answering "what sheets does this spreadsheet have".

Two confirmed production failures, from the same session (477b296a) and the
same workbook (data-file-5.xlsx — 10 sheets: Cover Sheet, Terms, Sheet1,
Working Sheet, Other, DropDown, Enquiry amendment form, Deviation List,
RFQ-Authorities, Delivery; verified by opening the file with openpyxl):

1. WRONG COUNT, no error anywhere. Asked "how many sheets", the answer was
   "17 sheets", listing "Cover Sheet_Raw", "Cover Sheet (block 2)",
   "Terms_Raw"... Those are ingestion artifacts, not sheets: one sheet yields
   a "_Raw" whole-sheet table plus a "(block N)" per detected sub-table.
   Both the stored FILE_META["toc"] and the code-exec sandbox's variable list
   are per-table, so every route to the answer counted tables. 'Delivery',
   which has no extractable table, was missing from all 17.

2. "CORRUPTED", four turns in a row. At 14:11 the model answered an export
   with an invented file_id, "data-file-5_9912" — a plausible shape that was
   never uploaded (grep of the full interaction log: it appears only in
   responses, never in an ingestion record, and no such file exists in
   uploads/). LangGraph checkpointed that answer, so from then on the model
   read its own invention back out of the conversation and passed it to
   tools. resolve_target_file returned it unchecked, get_all_real_tables
   found nothing under it, and run_code_on_files said only "No tables found
   for the selected file(s)" — which the model reported to the user as "the
   file could not be parsed. It may be corrupted... please re-upload".
   The same questions answered correctly in a fresh session the whole time,
   because a fresh session had no invented id to inherit.

The fixes: sheet names are recorded from the workbook at ingestion rather
than inferred from tables; a file_id must be real before a tool will act on
it; and an empty result says which of the two things it means.
"""
import os

import pandas as pd
import pytest

from app.graph import agent as A
from app.graph import tools as T


def _table(page_label, data):
    df = pd.DataFrame(data)
    df.attrs["page"] = page_label
    return df


@pytest.fixture
def workbook(monkeypatch):
    """A tabular file shaped like the real one: 3 sheets' worth of tables,
    but 4 sheets in the workbook — 'Delivery' has nothing extractable on it,
    exactly like the sheet the old answer dropped."""
    tables = [
        _table("Cover Sheet_Raw", {"col_0": ["QUEST FLOW CONTROLS LTD", "Gat No"],
                                   "col_1": [None, "Offer Date :"]}),
        _table("Cover Sheet (block 2)", {"Customer Name": ["Space for Name"]}),
        _table("Terms_Raw", {"col_0": ["Commercial Terms"]}),
        _table("Working Sheet", {"Category": ["Ball Valve"], "Rating": ["PN10"]}),
    ]
    from app import state as app_state
    # Patched in both places on purpose: agent.py imports this inside the
    # function body (so the helpers-module patch reaches it) while tools.py
    # binds it at module load (so it needs its own). Patching only helpers
    # silently gave every sheet "no table extracted".
    monkeypatch.setattr("app.tables.helpers.get_all_real_tables",
                        lambda fid, **kw: tables)
    monkeypatch.setattr(T, "get_all_real_tables", lambda fid, **kw: tables)
    monkeypatch.setitem(app_state.FILE_KIND, "df5", "tabular")
    monkeypatch.setitem(app_state.FILE_META, "df5", {
        "toc": [{"text": f"Sheet/Block: {t.attrs['page']}",
                 "page": t.attrs["page"]} for t in tables],
        "sheets": ["Cover Sheet", "Terms", "Working Sheet", "Delivery"],
    })
    monkeypatch.setattr(app_state, "FILE_ORIGINAL_NAME",
                        {"df5": "data-file-5.xlsx"})
    monkeypatch.setattr(app_state, "FILE_ORDER", ["df5"])
    return tables


# --- the sheet list itself ---------------------------------------------------

def test_sheet_names_come_from_the_workbook_not_the_tables(workbook):
    """4 tables spanning 3 sheets, in a 4-sheet workbook. The old answer was
    the table count."""
    assert A._tabular_sheet_names("df5") == [
        "Cover Sheet", "Terms", "Working Sheet", "Delivery"]


def test_a_sheet_with_no_extractable_table_is_still_a_sheet(workbook):
    """'Delivery' contributed no table, so every table-derived route dropped
    it — the user was told their 10-sheet file had sheets it doesn't, and
    not told about one it does."""
    listing = T._sheet_listing("df5")
    assert "Delivery" in listing
    assert "empty (no table extracted)" in listing


def test_the_count_is_the_real_sheet_count(workbook):
    assert "**4 sheet(s)**" in T._sheet_listing("df5")


def test_ingestion_artifacts_are_never_presented_as_sheet_names(workbook):
    listing = T._sheet_listing("df5")
    assert "_Raw" not in listing
    assert "(block 2)" not in listing


def test_row_count_comes_from_the_whole_sheet_not_a_sub_block(workbook):
    """Cover Sheet has a 2-row _Raw table and a 1-row block. The sheet has
    2 rows; reporting the block's 1 would understate it."""
    row = [r for r in T._sheet_listing("df5").splitlines()
           if r.startswith("| 1 |")][0]
    assert "| 2 |" in row  # rows column


def test_falls_back_to_table_labels_for_files_ingested_before_this(monkeypatch):
    """Files still resident in memory from an older boot have no "sheets"
    key. Deriving names from table labels is wrong-ish but far better than
    refusing to answer."""
    tables = [_table("Alpha_Raw", {"a": [1]}), _table("Alpha (block 2)", {"a": [1]}),
              _table("Beta", {"b": [2]})]
    monkeypatch.setattr("app.tables.helpers.get_all_real_tables",
                        lambda fid, **kw: tables)
    from app import state as app_state
    monkeypatch.setitem(app_state.FILE_KIND, "old", "tabular")
    monkeypatch.setitem(app_state.FILE_META, "old", {"toc": []})
    assert A._tabular_sheet_names("old") == ["Alpha", "Beta"]


# --- routing -----------------------------------------------------------------

def test_structure_question_about_a_spreadsheet_gets_sheets_not_blocks(workbook):
    """get_table_of_contents is where rule 5 sends every structure question.
    For a spreadsheet its stored toc is one entry per TABLE — that list is
    what became "17 sheets"."""
    out = T.get_table_of_contents.invoke({"file_id": "df5"})
    assert "**4 sheet(s)**" in out
    assert "Sheet/Block:" not in out


def test_list_sheets_refuses_politely_for_a_pdf(monkeypatch):
    from app import state as app_state
    monkeypatch.setitem(app_state.FILE_KIND, "pdf1", "docling")
    monkeypatch.setattr(app_state, "FILE_ORIGINAL_NAME", {"pdf1": "tender.pdf"})
    out = T.list_sheets.invoke({"file_id": "pdf1"})
    assert "not a spreadsheet" in out


@pytest.mark.parametrize("prompt", [
    "in data-file-5 how many sheet is there list out that names",
    "list out this file sheet names = data-file-5",
    "give me all sheet names list",
    "list every sheet name and how many rows each has in this file",
    "list out howmany sheet is there in data-file-5",
    "what tabs does this file have",
])
def test_every_phrasing_the_user_actually_typed_is_recognised(prompt):
    assert A._SHEET_ASK_RE.search(prompt)


@pytest.mark.parametrize("prompt", [
    "export sheet name = cover sheet into x.xlsx",
    "in data-file-5 check data like want to export sheet name = cover sheet "
    "in to new file sheet-01.xlsx",
])
def test_an_export_that_names_a_sheet_is_not_an_inventory_question(prompt):
    """These match the wording but ask for a FILE. Injecting "answer from
    this inventory" would talk the model out of the export."""
    assert A._SHEET_ASK_RE.search(prompt)          # wording matches...
    assert A._asks_for_a_file(prompt)              # ...but the guard excludes it


# --- the invented file_id ----------------------------------------------------

def test_an_invented_file_id_is_not_trusted(workbook, monkeypatch):
    """The whole second failure in one assertion: data-file-5_9912 was never
    uploaded, so acting on it can only produce a false 'corrupted' answer."""
    monkeypatch.setattr("app.persistence.get_file", lambda fid: None)
    monkeypatch.setattr("app.state.get_active_file_id", lambda: "df5")
    assert T.resolve_target_file("data-file-5_9912") == "df5"


def test_a_real_file_id_is_still_used_as_given(workbook):
    assert T.resolve_target_file("df5") == "df5"


def test_a_registered_but_not_yet_loaded_file_id_is_still_trusted(monkeypatch):
    """Restart empties memory; the registry is what survives it. Rejecting
    these would break every reference to a file uploaded before a restart."""
    monkeypatch.setattr("app.persistence.get_file",
                        lambda fid: {"file_id": fid, "path": "/tmp/x.xlsx",
                                     "original_filename": "x.xlsx"})
    assert T.resolve_target_file("cold_file") == "cold_file"


def test_unknown_file_id_never_reads_as_a_broken_document(monkeypatch):
    """The exact sentence the model turned into "may be corrupted, please
    re-upload" was a bare "No tables found for the selected file(s)."."""
    from app import state as app_state
    monkeypatch.setattr(app_state, "FILE_ORDER", ["df5"])
    monkeypatch.setattr(app_state, "FILE_KIND", {"df5": "tabular"})
    monkeypatch.setattr(app_state, "FILE_ORIGINAL_NAME", {"df5": "data-file-5.xlsx"})
    monkeypatch.setattr("app.tables.helpers.get_all_real_tables",
                        lambda fid, **kw: [])
    from app.query.code_exec import run_code_on_files
    text = run_code_on_files(["data-file-5_9912"], "how many sheets")["text"]
    assert "does not refer to any uploaded file" in text
    assert "data-file-5.xlsx (id=df5)" in text   # names what IS available
    assert "corrupt" not in text.lower()


# --- the invention outliving the bug it caused -------------------------------

def test_a_sentence_about_a_phantom_file_id_is_dropped(workbook, monkeypatch):
    """Confirmed live after the routing fix: the sheet list came back correct
    and the answer still ended "The file `data-file-5_9912` remains
    inaccessible and cannot be processed." — the invention outliving the bug
    that produced it, still frightening the user about a file they never
    had."""
    monkeypatch.setattr("app.persistence.get_file", lambda fid: None)
    text = ("Here are all 10 sheet names. The file `data-file-5_9912` "
            "remains inaccessible and cannot be processed.")
    out = A._strip_phantom_file_ids(text)
    assert "9912" not in out
    assert "Here are all 10 sheet names." in out


def test_a_sentence_about_a_real_file_id_is_kept(workbook):
    text = "Source: 'data-file-5.xlsx' (file_id `df5`) has 4 sheets."
    assert A._strip_phantom_file_ids(text) == text


def test_a_phantom_id_never_blanks_the_whole_answer(monkeypatch):
    """If the invention was the entire response, returning "" would leave an
    empty bubble — worse than the wrong sentence."""
    monkeypatch.setattr("app.persistence.get_file", lambda fid: None)
    text = "The file `data-file-9_9999` could not be read."
    assert A._strip_phantom_file_ids(text) == text


def test_a_filename_in_prose_is_not_mistaken_for_a_file_id(monkeypatch):
    """Only backticked file_id-shaped tokens count. A plain filename in a
    sentence must never cause that sentence to disappear."""
    monkeypatch.setattr("app.persistence.get_file", lambda fid: None)
    text = "Saved 27 rows to sheet-01.xlsx from data-file-5.xlsx."
    assert A._strip_phantom_file_ids(text) == text


# --- export of a sheet that exists but is empty ------------------------------

def test_an_empty_sheet_is_reported_as_empty_not_as_missing(workbook):
    """Reachable only now that sheet names come from the workbook: 'Delivery'
    is a real sheet with no table. "No sheet named 'Delivery'" would be a
    lie about a sheet the user can see in Excel."""
    out = A._deterministic_sheet_export("df5", "Delivery", "out.xlsx", "excel")
    assert "exists" in out and "empty" in out
    assert "No sheet named" not in out


def test_a_genuinely_absent_sheet_still_says_so(workbook):
    out = A._deterministic_sheet_export("df5", "Nonexistent", "out.xlsx", "excel")
    assert "No sheet named" in out


# --- against the real workbook ----------------------------------------------

REAL_XLSX = "uploads/data-file-5_3361_data-file-5.xlsx"


@pytest.mark.skipif(not os.path.exists(REAL_XLSX),
                    reason="real workbook not available")
def test_the_real_workbook_reports_its_ten_real_sheets(monkeypatch):
    """Ground truth, read straight from the file with openpyxl, against what
    ingestion now records. This is the assertion the shipped answer of
    "17 sheets" would have failed."""
    import openpyxl
    from app import state as app_state
    from app.ingestion.tabular_ingest import ingest_tabular

    monkeypatch.setattr("app.ingestion.tabular_ingest.embed_text",
                        lambda *a, **k: None)  # no vector store in tests
    monkeypatch.setattr(app_state, "TABULAR_TABLES", {})
    monkeypatch.setattr(app_state, "FILE_KIND", {})
    monkeypatch.setattr(app_state, "FILE_META", {})
    monkeypatch.setattr(app_state, "FILE_ORIGINAL_NAME", {"real5": "data-file-5.xlsx"})

    ingest_tabular(REAL_XLSX, "real5")
    expected = list(openpyxl.load_workbook(REAL_XLSX, read_only=True).sheetnames)

    assert A._tabular_sheet_names("real5") == expected
    assert len(expected) == 10
    listing = T._sheet_listing("real5")
    assert "**10 sheet(s)**" in listing
    for sheet in expected:
        assert sheet.strip() in listing
