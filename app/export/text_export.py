"""Routing prose (overview/TOC/answer) into whichever file format was requested."""
import re
from typing import Optional
import pandas as pd
from app.config import OUTPUT_DIR
import os


def clean_llm_csv(text: str) -> str:
    text = re.sub(r"^```(?:csv)?\s*", "", text.strip())
    text = re.sub(r"\s*```$", "", text.strip())
    return text.strip()


def _out_path(filename: str) -> str:
    return os.path.join(OUTPUT_DIR, filename)


def write_text_to_docx(text: str, filename: str, title: str = "Document") -> str:
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    for para in text.split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    path = _out_path(filename)
    doc.save(path)
    return f"Saved -> {path}"


def write_text_to_pptx(text: str, filename: str, title: str = "Document") -> str:
    from pptx import Presentation
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title
    body_slide = prs.slides.add_slide(prs.slide_layouts[1])
    tf = body_slide.placeholders[1].text_frame
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    tf.text = lines[0] if lines else ""
    for line in lines[1:20]:
        p = tf.add_paragraph()
        p.text = line
    path = _out_path(filename)
    prs.save(path)
    return f"Saved -> {path}"


def save_text_as(text: str, sink: Optional[str], filename: Optional[str],
                  default_stub: str, title: str) -> str:
    if sink == "docx":
        return write_text_to_docx(text, filename or f"{default_stub}.docx", title)
    if sink == "pptx":
        return write_text_to_pptx(text, filename or f"{default_stub}.pptx", title)
    if sink == "csv":
        out = _out_path(filename or f"{default_stub}.csv")
        pd.DataFrame({"content": text.split("\n")}).to_csv(out, index=False)
        return f"Saved as a single-column CSV (prose doesn't fit tabular data well) -> {out}"
    if sink == "excel":
        out = _out_path(filename or f"{default_stub}.xlsx")
        pd.DataFrame({"content": text.split("\n")}).to_excel(out, index=False)
        return f"Saved as a single-column sheet (prose doesn't fit tabular data well) -> {out}"
    if sink == "chart":
        return "A chart isn't a good fit for text content — try docx or pptx instead."
    return text
